import importlib.util
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path('/workspace/scratch/150b36589797/upload/01-CF-M02-Apostila-Didatica-1.docx')
SOURCE = Path('/workspace/scratch/150b36589797/upload/CF-M03-Direitos-sociais-nacionalidade-e-direitos-politicos-Teoria-completa.txt')
OUT = Path('/workspace/scratch/150b36589797/output/cf-m03/CF-M03-Apostila-Didatica-1.docx')


def load_base_builder():
    path = ROOT / 'scripts/build-cf-m01-docx.py'
    spec = importlib.util.spec_from_file_location('cf_m01_docx_builder', path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


B = load_base_builder()


def p(doc, text, style='Normal'):
    return B.add_rich_paragraph(doc, text, style)


def h(doc, text, level=1):
    return B.add_heading(doc, text, level)


def bullets(doc, items):
    for item in items:
        p(doc, item, 'List Bullet')


def table(doc, headers, rows):
    return B.add_table(doc, [headers, *rows])


def update_running_text(doc):
    for section in doc.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                for run in paragraph.runs:
                    run.text = run.text.replace('CF M02', 'CF M03')


def add_cover(doc):
    top = doc.add_paragraph(style='Normal')
    run = top.add_run('BASE COMPLETA')
    B.add_run_format(run, size=12, bold=True, color='142F50')
    top.paragraph_format.space_after = Pt(18)

    doc.add_paragraph('Direitos sociais nacionalidade e direitos políticos', style='Title')
    doc.add_paragraph('Direito Constitucional • CF M03', style='Subtitle')
    doc.add_paragraph('Teoria explicada e aprofundamento para concursos de tribunais', style='Subtitle')
    meta = doc.add_paragraph('Edição didática • 22 de outubro de 2026', style='Normal')
    meta.paragraph_format.space_before = Pt(10)
    intro = doc.add_paragraph(
        'Aprenda a distinguir direitos sociais, nacionalidade, participação política e partidos por meio de regras, exceções e casos concretos.',
        style='Normal',
    )
    intro.paragraph_format.space_before = Pt(18)
    intro.paragraph_format.space_after = Pt(20)


QUESTIONS = [
    {
        'n': 1, 'difficulty': 'Fácil',
        'q': 'Uma empresa possui exatamente 200 empregados. A Constituição exige a eleição do representante previsto no art. 11?',
        'options': ['Sim, porque o limite é de 200 ou mais.', 'Não, porque o texto exige mais de 200 empregados.', 'Sim, desde que haja sindicato.', 'Não, porque essa representação só existe no serviço público.'],
        'answer': 1,
        'explanation': 'O art. 11 utiliza a expressão “mais de duzentos empregados”. Exatamente 200 não satisfaz o requisito numérico.'
    },
    {
        'n': 2, 'difficulty': 'Fácil',
        'q': 'Qual pessoa pode alistar-se e votar facultativamente, mas não pode ser eleita?',
        'options': ['Estrangeiro residente.', 'Conscrito durante o serviço militar obrigatório.', 'Analfabeto.', 'Menor de 16 anos.'],
        'answer': 2,
        'explanation': 'O analfabeto é alistável com voto facultativo, porém inelegível. Estrangeiros e conscritos são inalistáveis.'
    },
    {
        'n': 3, 'difficulty': 'Fácil',
        'q': 'Após a EC 131/2023, a aquisição de outra nacionalidade por brasileiro provoca automaticamente a perda da nacionalidade brasileira?',
        'options': ['Sim, em qualquer hipótese.', 'Sim, salvo autorização judicial.', 'Não. A mera aquisição deixou de ser causa automática de perda.', 'Não, mas apenas para brasileiro naturalizado.'],
        'answer': 2,
        'explanation': 'A EC 131/2023 suprimiu a perda pela mera aquisição de outra nacionalidade. O regime atual prevê cancelamento judicial da naturalização nas hipóteses constitucionais e pedido expresso de perda, ressalvada a apatridia.'
    },
    {
        'n': 4, 'difficulty': 'Fácil',
        'q': 'Qual cargo possui idade mínima constitucional de 18 anos?',
        'options': ['Vereador.', 'Prefeito.', 'Deputado Estadual.', 'Senador.'],
        'answer': 0,
        'explanation': 'A escala constitucional é 35, 30, 21 e 18. Vereador é o único cargo da lista com idade mínima de 18 anos.'
    },
    {
        'n': 5, 'difficulty': 'Média',
        'q': 'Uma brasileira nasceu no exterior, filha de mãe brasileira que não estava a serviço do Brasil, e foi registrada em repartição brasileira competente. Qual é sua situação?',
        'options': ['Estrangeira, até residir no Brasil por quinze anos.', 'Brasileira naturalizada.', 'Brasileira nata.', 'Brasileira equiparada.'],
        'answer': 2,
        'explanation': 'O registro em repartição brasileira competente é hipótese autônoma de nacionalidade originária para filho de pai brasileiro ou mãe brasileira nascido no exterior.'
    },
    {
        'n': 6, 'difficulty': 'Média',
        'q': 'Prefeito em primeiro mandato pretende concorrer ao cargo de governador. Qual providência constitucional é necessária?',
        'options': ['Afastar-se três meses antes.', 'Renunciar até seis meses antes do pleito.', 'Permanecer no cargo, porque ainda não foi reeleito.', 'Obter autorização da Câmara Municipal.'],
        'answer': 1,
        'explanation': 'A dispensa de renúncia vale para a reeleição ao mesmo cargo. Para concorrer a outro cargo, o chefe do Executivo deve renunciar até seis meses antes do pleito.'
    },
    {
        'n': 7, 'difficulty': 'Média',
        'q': 'Em relação à criação de partido político, qual sequência está correta?',
        'options': ['Registro no TSE e depois personalidade civil.', 'Autorização do Congresso e depois registro civil.', 'Personalidade jurídica na forma da lei civil e depois registro do estatuto no TSE.', 'Registro no STF e depois autorização do TSE.'],
        'answer': 2,
        'explanation': 'A Constituição separa duas etapas: aquisição da personalidade jurídica conforme a lei civil e posterior registro do estatuto no Tribunal Superior Eleitoral.'
    },
    {
        'n': 8, 'difficulty': 'Média',
        'q': 'A contribuição assistencial definida em acordo ou convenção coletiva pode alcançar empregados não sindicalizados?',
        'options': ['Nunca, porque isso equivaleria à filiação obrigatória.', 'Sim, desde que seja assegurado o direito de oposição.', 'Sim, sem possibilidade de oposição.', 'Somente mediante autorização judicial individual.'],
        'answer': 1,
        'explanation': 'Segundo o Tema 935 do STF, a contribuição assistencial pode abranger toda a categoria, inclusive não sindicalizados, desde que exista direito de oposição.'
    },
    {
        'n': 9, 'difficulty': 'Difícil',
        'q': 'Um brasileiro naturalizado é acusado de crime comum cometido depois da naturalização, sem relação com tráfico de drogas. A extradição é constitucionalmente possível?',
        'options': ['Sim, porque naturalizado sempre pode ser extraditado.', 'Não por esse fundamento, pois o crime comum deve ser anterior à naturalização; permanece a hipótese constitucional relativa ao tráfico.', 'Sim, desde que o delito seja punido com reclusão.', 'Não, porque brasileiro naturalizado nunca pode ser extraditado.'],
        'answer': 1,
        'explanation': 'O naturalizado pode ser extraditado por crime comum praticado antes da naturalização ou por comprovado envolvimento com tráfico ilícito de entorpecentes e drogas afins, na forma constitucional.'
    },
    {
        'n': 10, 'difficulty': 'Difícil',
        'q': 'Lei publicada oito meses antes da eleição altera o processo eleitoral. Qual afirmação expressa corretamente o art. 16?',
        'options': ['A lei não entra em vigor antes de um ano.', 'A lei entra em vigor na publicação, mas não se aplica à eleição que ocorra até um ano da vigência.', 'A lei aplica-se imediatamente porque foi publicada antes da campanha.', 'A lei depende de resolução do TSE para entrar em vigor.'],
        'answer': 1,
        'explanation': 'O art. 16 distingue vigência e aplicação: a lei vigora desde a publicação, mas não incide na eleição realizada dentro do intervalo constitucional de um ano.'
    },
]


REVIEW_QUESTIONS = [
    {
        'n': 1, 'difficulty': 'Fácil',
        'q': 'Certo ou errado: o voto obrigatório integra a cláusula pétrea do art. 60, § 4º, II.',
        'options': ['Certo.', 'Errado.'], 'answer': 1,
        'explanation': 'A proteção alcança o voto direto, secreto, universal e periódico. A obrigatoriedade não aparece nessa enumeração.'
    },
    {
        'n': 2, 'difficulty': 'Fácil',
        'q': 'Certo ou errado: brasileiro nato nunca será extraditado.',
        'options': ['Certo.', 'Errado.'], 'answer': 0,
        'explanation': 'A vedação é absoluta para o brasileiro nato. As hipóteses excepcionais de extradição referem-se ao naturalizado.'
    },
    {
        'n': 3, 'difficulty': 'Média',
        'q': 'Certo ou errado: coligações partidárias são constitucionalmente proibidas tanto em eleições proporcionais quanto majoritárias.',
        'options': ['Certo.', 'Errado.'], 'answer': 1,
        'explanation': 'A vedação constitucional recai sobre eleições proporcionais. O regime majoritário admite coligações conforme as regras aplicáveis.'
    },
    {
        'n': 4, 'difficulty': 'Média',
        'q': 'Certo ou errado: a condenação criminal transitada em julgado suspende direitos políticos somente quando a pena é privativa de liberdade.',
        'options': ['Certo.', 'Errado.'], 'answer': 1,
        'explanation': 'Segundo o Tema 370 do STF, a suspensão incide enquanto durarem os efeitos da condenação, independentemente da espécie de pena.'
    },
    {
        'n': 5, 'difficulty': 'Difícil',
        'q': 'Certo ou errado: dissolvido o vínculo conjugal durante o mandato, desaparece automaticamente a inelegibilidade reflexa do ex-cônjuge.',
        'options': ['Certo.', 'Errado.'], 'answer': 1,
        'explanation': 'A Súmula Vinculante 18 impede essa conclusão: a dissolução ocorrida no curso do mandato não afasta a inelegibilidade do art. 14, § 7º.'
    },
]


def add_question(doc, q, review=False):
    label = 'Revisão' if review else 'Questão'
    h(doc, f"{label} {q['n']} {q['difficulty']}", 2)
    p(doc, q['q'])
    for i, option in enumerate(q['options']):
        p(doc, f"{chr(65 + i)}) {option}")
    p(doc, f"Resposta comentada. {chr(65 + q['answer'])}. {q['explanation']}")


def build_body(doc):
    h(doc, 'Leitura orientada', 1)
    p(doc, 'Imagine quatro situações. Uma trabalhadora quer saber quais proteções mínimas a Constituição lhe garante. Uma pessoa nascida no exterior pergunta se é brasileira nata. Um prefeito pretende disputar o governo do Estado. Um partido quer registrar seu estatuto e acessar recursos públicos. Embora pareçam problemas separados, todos exigem o mesmo método: identificar a regra constitucional, localizar a exceção e descobrir qual consequência jurídica muda o caso.')
    p(doc, 'Se estiver começando, estude os quatro blocos na ordem desta apostila. Em cada um, primeiro compreenda a lógica; depois memorize listas, idades e prazos. Se já possui base, use os casos resolvidos e as questões para testar se consegue justificar a resposta sem depender apenas de palavras-chave.')
    p(doc, 'O módulo considera a Constituição atualizada até 22 de outubro de 2026, inclusive as alterações recentes sobre jornada de trabalho e repouso semanal promovidas pela EC 141/2026. Em prova que indique data anterior, observe o marco temporal do enunciado.')

    h(doc, 'Percurso do módulo', 2)
    bullets(doc, [
        'Aula 1 • Direitos sociais e a lógica da igualdade material.',
        'Aula 2 • Direitos dos trabalhadores e números que mudam o gabarito.',
        'Aula 3 • Liberdade sindical, greve e representação no local de trabalho.',
        'Aula 4 • Nacionalidade originária, naturalização e portugueses equiparados.',
        'Aula 5 • Cargos privativos, perda da nacionalidade e extradição.',
        'Aula 6 • Capacidade eleitoral ativa e passiva.',
        'Aula 7 • Elegibilidade, inelegibilidades, prazos e militares.',
        'Aula 8 • Partidos políticos, fidelidade e financiamento.',
        'Aula 9 • Lei seca e jurisprudência aplicadas.',
        'Aula 10 • Questões, revisão ativa e checklist de domínio.',
    ])

    h(doc, '1 Direitos sociais e igualdade material', 1)
    h(doc, '1.1 O que os direitos sociais fazem', 2)
    p(doc, 'Direitos sociais não são apenas nomes reunidos em uma lista. Eles procuram reduzir desigualdades reais e criar condições mínimas para que a liberdade possa ser exercida. Por isso, alguns exigem prestações do Estado, como políticas de saúde e educação; outros protegem relações de trabalho; e outros funcionam como liberdades coletivas, a exemplo da associação sindical e da greve.')
    p(doc, 'Caso explicado. Dizer que uma pessoa é livre para estudar pouco resolve se ela não possui acesso à escola. A igualdade material pergunta quais condições concretas permitem que a liberdade exista. Essa é a chave para entender por que direitos sociais integram o catálogo de direitos fundamentais.')

    h(doc, '1.2 Artigo 6º e o núcleo geral', 2)
    p(doc, 'O art. 6º reúne educação, saúde, alimentação, trabalho, moradia, transporte, lazer, segurança, previdência social, proteção à maternidade e à infância e assistência aos desamparados. O rol orienta o estudo, mas não esgota todos os direitos sociais existentes na Constituição. Direitos trabalhistas, liberdade sindical e greve também pertencem ao tema, embora apareçam nos artigos seguintes.')
    table(doc, ['Direito', 'Como compreender', 'Exemplo de incidência'], [
        ['Educação e saúde', 'Prestação pública e garantia de acesso', 'Escola pública e atendimento pelo sistema de saúde'],
        ['Trabalho e previdência', 'Proteção da vida profissional e dos riscos sociais', 'Seguro-desemprego e benefícios previdenciários'],
        ['Moradia e transporte', 'Condições materiais de vida e mobilidade', 'Políticas habitacionais e transporte acessível'],
        ['Maternidade infância e assistência', 'Proteção de situações de vulnerabilidade', 'Benefício assistencial e políticas de proteção'],
    ])
    p(doc, 'Aprofundamento. O parágrafo único do art. 6º assegura a brasileiro em situação de vulnerabilidade social uma renda básica familiar por programa permanente de transferência de renda, nos termos da lei e respeitadas as regras fiscais e orçamentárias. A Constituição reconhece o direito, mas remete sua concretização ao desenho legal do programa.')

    h(doc, '1.3 Artigo 7º e o piso constitucional do trabalhador', 2)
    p(doc, 'O art. 7º estabelece um piso de proteção para trabalhadores urbanos e rurais. A expressão “além de outros que visem à melhoria de sua condição social” mostra que a lista constitucional não impede proteção adicional. Em questões objetivas, a banca costuma trocar números, retirar exceções ou transformar uma regra dependente de lei em direito de aplicação idêntica em qualquer situação.')
    table(doc, ['Tema', 'Regra constitucional', 'Ponto de atenção'], [
        ['Emprego', 'Proteção contra despedida arbitrária ou sem justa causa nos termos de lei complementar', 'A Constituição prevê proteção e indenização compensatória, mas remete detalhes à lei'],
        ['Salário', 'Mínimo nacionalmente unificado e irredutibilidade, salvo negociação coletiva', 'Não trate a irredutibilidade como absoluta'],
        ['Jornada', 'Até 8 horas diárias e 40 semanais no texto permanente', 'A EC 141/2026 estabeleceu transição para o limite semanal'],
        ['Turnos de revezamento', '6 horas, salvo negociação coletiva', 'A exceção depende de negociação coletiva'],
        ['Serviço extraordinário', 'Adicional de pelo menos 50%', 'A Constituição fixa o mínimo'],
        ['Férias', 'Remuneração acrescida de pelo menos um terço', 'O adicional incide sobre a remuneração de férias'],
        ['Licença gestante', '120 dias sem prejuízo do emprego e do salário', 'Não confundir com prazos legais ampliados em regimes específicos'],
        ['Aviso prévio', 'Proporcional ao tempo de serviço, mínimo de 30 dias', 'A Constituição fixa o mínimo'],
        ['Prescrição', 'Créditos dos últimos 5 anos, ação até 2 anos após a extinção', 'Use a fórmula 5 mais 2'],
        ['Trabalho do menor', 'Regra geral a partir de 16; aprendiz desde 14', 'Menor de 18 não pode exercer trabalho noturno, perigoso ou insalubre'],
        ['Trabalhador avulso', 'Igualdade de direitos com trabalhador de vínculo permanente', 'Avulso não significa empregado permanente'],
    ])
    p(doc, 'Atualização jurídica. A EC 141/2026 passou a prever 40 horas semanais no texto permanente e uma transição constitucional: manutenção de 44 horas até o término dos dois primeiros meses da publicação; 42 horas pelo período seguinte definido na emenda; e, ao final, 40 horas. A emenda também instituiu dois dias de repouso semanal remunerado após o período inicial, sendo um deles preferencialmente aos domingos. Como a alteração é muito recente, confira a data da questão e a regra transitória aplicável.')
    p(doc, 'Caso resolvido. Se uma questão de prova atual disser apenas que a Constituição fixa jornada máxima de 44 horas semanais, a assertiva está desatualizada. Se o enunciado tratar do período de transição, é preciso identificar a data: texto permanente e aplicação temporal não são a mesma coisa.')
    p(doc, 'Empregados domésticos recebem diversos direitos do art. 7º. A EC 72/2013 ampliou a proteção, mas o próprio texto constitucional submete alguns direitos às condições previstas em lei. Por isso, a frase “todos os incisos aplicam-se indistintamente” é ampla demais.')

    h(doc, '1.4 Liberdade sindical', 2)
    p(doc, 'A liberdade sindical protege a criação, a atuação e a autonomia das entidades representativas, sem transformar a filiação em obrigação. O registro no órgão competente não equivale a autorização estatal: serve para organizar o sistema e verificar a unicidade na base territorial.')
    table(doc, ['Regra', 'O que significa', 'Erro comum'], [
        ['Sem autorização estatal', 'A lei não pode exigir permissão para fundar sindicato, ressalvado o registro', 'Confundir registro com autorização política'],
        ['Unicidade sindical', 'Uma organização da mesma categoria por base territorial', 'Admitir duas entidades idênticas na mesma base'],
        ['Base mínima', 'Não pode ser inferior à área de um Município', 'Aceitar base de bairro'],
        ['Defesa da categoria', 'O sindicato atua judicial e administrativamente', 'Limitar a defesa apenas aos filiados'],
        ['Filiação livre', 'Ninguém é obrigado a filiar-se ou permanecer filiado', 'Confundir contribuição com filiação'],
        ['Negociação coletiva', 'Participação sindical obrigatória', 'Dispensar o sindicato na negociação coletiva'],
        ['Dirigente sindical', 'Proteção desde o registro da candidatura e, se eleito, até um ano após o mandato', 'Tratar a garantia como absoluta diante de falta grave'],
    ])
    p(doc, 'Jurisprudência aplicada. No Tema 935, o STF considerou constitucional a contribuição assistencial instituída por acordo ou convenção coletiva para todos os empregados da categoria, inclusive não sindicalizados, desde que seja assegurado o direito de oposição. Cobrança coletiva e filiação obrigatória são problemas diferentes.')

    h(doc, '1.5 Greve participação e representação', 2)
    p(doc, 'A greve é direito dos trabalhadores, que decidem sobre a oportunidade de exercê-la e os interesses a defender. A lei disciplina serviços essenciais e o atendimento das necessidades inadiáveis da comunidade. O abuso não elimina o direito, mas sujeita os responsáveis às consequências legais.')
    p(doc, 'Para agentes públicos, não aplique automaticamente a mesma disciplina. O STF considera incompatível com a Constituição o exercício de greve por policiais civis e por servidores que atuem diretamente na segurança pública. O fundamento está ligado à continuidade de função estatal essencial e ao regime constitucional da segurança.')
    p(doc, 'O art. 10 assegura participação de trabalhadores e empregadores em colegiados de órgãos públicos quando interesses profissionais ou previdenciários sejam discutidos. O art. 11 garante representante para promover entendimento direto com o empregador apenas em empresas com mais de 200 empregados. A palavra “mais” decide o gabarito.')

    h(doc, '2 Nacionalidade', 1)
    h(doc, '2.1 Como organizar o tema', 2)
    p(doc, 'Nacionalidade é o vínculo jurídico-político entre a pessoa e o Estado. A nacionalidade originária resulta das hipóteses constitucionais de nascimento e forma o grupo dos brasileiros natos. A derivada decorre da naturalização. A regra é a igualdade entre natos e naturalizados; diferenças só podem existir quando a própria Constituição as estabelece.')
    p(doc, 'Método de prova. Em caso de nascimento, responda quatro perguntas: onde a pessoa nasceu; qual a nacionalidade dos pais; algum deles estava a serviço do Brasil ou de país estrangeiro; houve registro em repartição brasileira ou residência e opção após a maioridade?')

    h(doc, '2.2 Brasileiros natos', 2)
    table(doc, ['Situação', 'Requisito constitucional', 'Resultado'], [
        ['Nascimento no Brasil', 'Pais estrangeiros não podem estar a serviço do país de origem', 'Brasileiro nato pelo critério territorial'],
        ['Nascimento no exterior com serviço do Brasil', 'Pai brasileiro ou mãe brasileira a serviço da República Federativa do Brasil', 'Brasileiro nato'],
        ['Nascimento no exterior com registro', 'Pai brasileiro ou mãe brasileira e registro em repartição brasileira competente', 'Brasileiro nato'],
        ['Nascimento no exterior com residência e opção', 'Pai brasileiro ou mãe brasileira; residência no Brasil; opção depois da maioridade', 'Brasileiro nato após o cumprimento dos requisitos'],
    ])
    p(doc, 'Caso resolvido. Ana nasceu no Canadá, filha de mãe brasileira que não estava a serviço do Brasil, e foi registrada no consulado. Ela é brasileira nata. O registro é hipótese autônoma; não é necessário somá-lo à residência no Brasil e à opção.')
    p(doc, 'Pegadinha. Se a criança nasce no Brasil e os pais são estrangeiros, a regra continua sendo a nacionalidade brasileira. A exceção exige que eles estejam a serviço do país estrangeiro. Ser estrangeiro, isoladamente, não basta.')

    h(doc, '2.3 Naturalizados e portugueses equiparados', 2)
    p(doc, 'A naturalização ordinária constitucional para originários de países de língua portuguesa exige residência por um ano ininterrupto e idoneidade moral. A naturalização extraordinária alcança estrangeiro de qualquer nacionalidade residente no Brasil há mais de quinze anos ininterruptos, sem condenação penal, desde que requeira a nacionalidade.')
    p(doc, 'O português com residência permanente no Brasil pode receber direitos inerentes ao brasileiro se houver reciprocidade em favor de brasileiros. Isso é equiparação de direitos. Não transforma automaticamente o português em brasileiro nato ou naturalizado.')

    h(doc, '2.4 Cargos privativos de brasileiro nato', 2)
    table(doc, ['Cargo ou função', 'Por que memorizar'], [
        ['Presidente e Vice-Presidente da República', 'Chefes do Executivo federal'],
        ['Presidente da Câmara dos Deputados', 'Linha sucessória'],
        ['Presidente do Senado Federal', 'Linha sucessória'],
        ['Ministro do Supremo Tribunal Federal', 'Cúpula do Judiciário constitucional'],
        ['Carreira diplomática', 'Representação internacional do Estado'],
        ['Oficial das Forças Armadas', 'Comando militar'],
        ['Ministro de Estado da Defesa', 'Direção civil da Defesa'],
    ])
    p(doc, 'A lista do art. 12, § 3º é taxativa. A Constituição também exige brasileiros natos entre os cidadãos que integram o Conselho da República, mas isso não transforma essa função em um oitavo cargo da lista do § 3º.')

    h(doc, '2.5 Perda reaquisição e extradição', 2)
    p(doc, 'A EC 131/2023 alterou a lógica da perda da nacionalidade. A mera aquisição de outra nacionalidade não provoca mais perda automática da brasileira. O texto atual prevê cancelamento judicial da naturalização por fraude relacionada ao processo ou por atentado contra a ordem constitucional e o Estado Democrático, além do pedido expresso de perda perante autoridade brasileira competente, ressalvadas situações de apatridia.')
    p(doc, 'A renúncia não impede posterior reaquisição da nacionalidade brasileira originária, nos termos da lei. Em prova, diferencie três ideias: possuir duas nacionalidades; pedir expressamente a perda; e ter a naturalização cancelada judicialmente. Elas não produzem o mesmo resultado.')
    table(doc, ['Pessoa', 'Regra de extradição'], [
        ['Brasileiro nato', 'Nunca será extraditado'],
        ['Brasileiro naturalizado', 'Crime comum anterior à naturalização ou comprovado envolvimento com tráfico ilícito de entorpecentes e drogas afins'],
        ['Estrangeiro', 'Em regra pode ser extraditado, ressalvadas vedações como crime político ou de opinião'],
    ])
    p(doc, 'O art. 13 completa o bloco: português é o idioma oficial; bandeira, hino, armas e selo são símbolos nacionais; Estados, Distrito Federal e Municípios podem possuir símbolos próprios.')

    h(doc, '3 Direitos políticos', 1)
    h(doc, '3.1 Sufrágio voto e capacidades eleitorais', 2)
    p(doc, 'Sufrágio é o direito de participação política e abrange votar e ser votado. Voto é o instrumento pelo qual se exerce a capacidade eleitoral ativa. Elegibilidade corresponde à capacidade passiva. Essa separação resolve boa parte das questões: o analfabeto possui capacidade ativa facultativa, mas não capacidade passiva.')
    p(doc, 'A soberania popular é exercida pelo sufrágio universal e pelo voto direto e secreto, com valor igual para todos, além de plebiscito, referendo e iniciativa popular. O voto direto, secreto, universal e periódico é cláusula pétrea. A obrigatoriedade do voto não integra essa enumeração.')
    table(doc, ['Situação', 'Alistamento e voto', 'Consequência'], [
        ['Maior de 18 e menor de 70', 'Obrigatórios, salvo situação constitucional específica', 'Capacidade ativa regular'],
        ['Analfabeto', 'Facultativos', 'Pode votar, mas é inelegível'],
        ['Maior de 70', 'Facultativos', 'Pode escolher se vota'],
        ['Maior de 16 e menor de 18', 'Facultativos', 'Ainda não possui idade para cargo eletivo'],
        ['Estrangeiro', 'Inalistável', 'Não vota nem é elegível'],
        ['Conscrito no serviço militar obrigatório', 'Inalistável durante o serviço', 'Não exerce capacidade eleitoral nesse período'],
    ])

    h(doc, '3.2 Condições de elegibilidade e idades', 2)
    table(doc, ['Condição', 'Como verificar'], [
        ['Nacionalidade brasileira', 'Nato ou naturalizado; cargo privativo exige nato'],
        ['Pleno exercício dos direitos políticos', 'A suspensão impede o preenchimento da condição'],
        ['Alistamento eleitoral', 'Sem alistamento não há capacidade passiva'],
        ['Domicílio eleitoral', 'Deve existir na circunscrição, conforme disciplina legal'],
        ['Filiação partidária', 'O regime atual não admite candidatura avulsa'],
        ['Idade mínima', '35 Presidente Vice e Senador; 30 Governador e Vice; 21 Deputados Prefeito Vice-Prefeito e juiz de paz; 18 Vereador'],
    ])
    p(doc, 'Técnica de memorização. Leia a sequência 35, 30, 21, 18 de cima para baixo. Agrupe os cargos: Presidência e Senado; governo estadual; cargos legislativos e executivos locais de maior responsabilidade; vereador. A idade é verificada no momento definido pela legislação eleitoral, preservadas as regras constitucionais.')

    h(doc, '3.3 Inelegibilidades reeleição e desincompatibilização', 2)
    p(doc, 'Inalistáveis e analfabetos são inelegíveis. Para chefes do Executivo, separe duas perguntas. Se a pessoa busca o mesmo cargo, pode haver uma única reeleição para período subsequente. Se pretende disputar outro cargo, deve renunciar ao mandato até seis meses antes do pleito.')
    p(doc, 'A inelegibilidade reflexa alcança, no território de jurisdição do titular, cônjuge e parentes consanguíneos ou afins até o segundo grau ou por adoção do Presidente, Governador, Prefeito ou de quem os tenha substituído nos seis meses anteriores. A exceção constitucional protege quem já é titular de mandato eletivo e concorre à reeleição. Pela Súmula Vinculante 18, a dissolução conjugal durante o mandato não afasta a inelegibilidade.')
    p(doc, 'Caso de comparação. Prefeito candidato à reeleição não precisa renunciar seis meses antes. Prefeito candidato a governador precisa renunciar. A pessoa é a mesma, mas o destino da candidatura altera a regra.')
    p(doc, 'O militar alistável é elegível. Com menos de dez anos de serviço, deve afastar-se da atividade. Com mais de dez anos, será agregado e, se eleito, passará automaticamente para a inatividade no ato da diplomação.')

    h(doc, '3.4 AIME consultas locais e art. 15', 2)
    p(doc, 'A ação de impugnação de mandato eletivo pode ser proposta perante a Justiça Eleitoral em quinze dias contados da diplomação, com provas de abuso do poder econômico, corrupção ou fraude. Tramita em segredo de justiça, e o autor responde se agir temerariamente ou com manifesta má-fé.')
    p(doc, 'Consultas populares sobre questões locais aprovadas pelas Câmaras Municipais podem ocorrer com as eleições municipais se forem encaminhadas à Justiça Eleitoral até noventa dias antes do pleito, respeitados os limites operacionais.')
    p(doc, 'A Constituição veda a cassação dos direitos políticos. O art. 15 admite perda ou suspensão nas hipóteses enumeradas: cancelamento da naturalização por sentença transitada em julgado; incapacidade civil absoluta; condenação criminal transitada em julgado enquanto durarem seus efeitos; recusa de cumprir obrigação a todos imposta ou prestação alternativa; e improbidade administrativa nos termos constitucionais. O artigo não rotula cada inciso como perda ou suspensão, razão pela qual classificações doutrinárias devem ser usadas com cautela.')
    p(doc, 'Jurisprudência aplicada. O Tema 370 do STF afirma que a suspensão por condenação criminal definitiva independe da espécie de pena. O Tema 1.190 esclarece que essa suspensão, por si só, não impede nomeação e posse em concurso público, desde que observadas a compatibilidade com a infração e as condições de exercício fixadas pelo Tribunal.')

    h(doc, '3.5 Anterioridade eleitoral', 2)
    p(doc, 'A lei que altera o processo eleitoral entra em vigor na data da publicação, mas não se aplica à eleição que ocorra até um ano da sua vigência. Vigência e aplicabilidade eleitoral são planos diferentes. A regra protege segurança jurídica, igualdade de chances e confiança no processo eleitoral.')

    h(doc, '4 Partidos políticos', 1)
    h(doc, '4.1 Liberdade autonomia e limites', 2)
    p(doc, 'A Constituição assegura liberdade de criação, fusão, incorporação e extinção de partidos, respeitados soberania nacional, regime democrático, pluripartidarismo e direitos fundamentais. Os partidos devem possuir caráter nacional, prestar contas à Justiça Eleitoral e não podem receber recursos de entidade ou governo estrangeiro nem subordinar-se a eles.')
    p(doc, 'A autonomia partidária alcança estrutura interna, organização, funcionamento e escolha de coligações nas eleições majoritárias. Coligações são vedadas nas eleições proporcionais. Também não existe obrigatoriedade de vinculação entre candidaturas nacionais, estaduais, distritais e municipais.')
    table(doc, ['Etapa ou tema', 'Regra correta', 'Troca comum em prova'], [
        ['Personalidade', 'Adquirida na forma da lei civil', 'Afirmar que nasce com o registro no TSE'],
        ['Registro eleitoral', 'Estatuto registrado no TSE depois da personalidade', 'Inverter a sequência'],
        ['Coligações', 'Admitidas no majoritário e vedadas no proporcional', 'Proibir em qualquer eleição'],
        ['Disciplina e fidelidade', 'Devem constar do estatuto', 'Tratar o vínculo partidário como irrelevante'],
        ['Organização paramilitar', 'Expressamente vedada', 'Admiti-la como manifestação da autonomia'],
    ])

    h(doc, '4.2 Desempenho fidelidade e financiamento', 2)
    p(doc, 'O acesso ao fundo partidário e à propaganda gratuita depende do cumprimento alternativo dos requisitos constitucionais de desempenho. No estágio final da cláusula, o partido deve obter três por cento dos votos válidos para a Câmara dos Deputados, distribuídos em pelo menos um terço das unidades da Federação, com mínimo de dois por cento em cada uma delas, ou eleger quinze Deputados Federais distribuídos em pelo menos um terço das unidades da Federação.')
    p(doc, 'O eleito por partido que não alcance a cláusula preserva o mandato e pode filiar-se, sem perda do mandato, a partido que a tenha alcançado, sem que a migração conte para distribuição de recursos e tempo gratuito. Deputados e Vereadores que deixam o partido pelo qual foram eleitos perdem o mandato, salvo anuência partidária ou outra hipótese legal de justa causa.')
    p(doc, 'A Constituição determina aplicação mínima de cinco por cento do fundo partidário em programas de promoção da participação política das mulheres. Também reserva mínimo de trinta por cento dos recursos de campanha e do tempo de propaganda às candidatas, proporcionalmente ao número de candidaturas. A EC 133/2024 acrescentou regra de aplicação de trinta por cento dos recursos destinados às campanhas em candidaturas de pessoas pretas e pardas, nas circunscrições escolhidas segundo os interesses e estratégias partidárias.')

    h(doc, '5 Método integrado para resolver casos', 1)
    table(doc, ['Se o enunciado tratar de', 'Pergunta de controle', 'Palavra decisiva'], [
        ['Direito social', 'É prestação, proteção trabalhista ou liberdade coletiva?', 'lei, mínimo, salvo, negociação'],
        ['Nacionalidade', 'Onde nasceu e qual vínculo dos pais com o Brasil?', 'registro, serviço, residência, opção'],
        ['Candidatura', 'A pessoa pode votar e também pode ser votada?', 'idade, filiação, outro cargo, seis meses'],
        ['Partido político', 'A questão trata de personalidade civil ou registro eleitoral?', 'depois, proporcional, desempenho'],
    ])
    bullets(doc, [
        'Primeiro identifique o artigo ou bloco constitucional.',
        'Depois sublinhe expressões quantitativas: mais de, até, um ano, quinze dias e noventa dias.',
        'Separe regra permanente de norma transitória, principalmente após a EC 141/2026.',
        'Teste a exceção com um caso concreto antes de marcar a alternativa.',
        'Em jurisprudência, use a tese apenas para o problema que ela efetivamente resolveu.',
    ])

    h(doc, '6 Lei seca essencial', 1)
    table(doc, ['Dispositivo', 'Núcleo para revisão'], [
        ['Art. 6º', 'Rol dos direitos sociais e renda básica familiar para brasileiro em vulnerabilidade'],
        ['Art. 7º', 'Piso trabalhista, jornada, repouso, adicionais, licenças, prescrição e trabalho do menor'],
        ['Art. 8º', 'Liberdade e unicidade sindical, base territorial, representação e dirigente'],
        ['Arts. 9º a 11', 'Greve, participação em colegiados e representante em empresa com mais de 200 empregados'],
        ['Arts. 12 e 13', 'Natos, naturalizados, cargos privativos, perda, extradição, idioma e símbolos'],
        ['Art. 14', 'Voto, alistamento, elegibilidade, idades, reeleição, renúncia, AIME e consultas locais'],
        ['Arts. 15 e 16', 'Vedação à cassação, perda ou suspensão e anterioridade eleitoral'],
        ['Art. 17', 'Criação, autonomia, registro, coligações, desempenho, fidelidade e financiamento'],
    ])

    h(doc, '7 Jurisprudência aplicada', 1)
    table(doc, ['Tema', 'Entendimento que muda a resposta'], [
        ['STF Tema 935', 'Contribuição assistencial para toda a categoria é possível com direito de oposição'],
        ['STF Tema 541', 'Greve é vedada a policiais civis e servidores diretamente ligados à segurança pública'],
        ['STF Tema 370', 'Suspensão por condenação definitiva independe da espécie de pena'],
        ['STF Tema 1.190', 'Suspensão não impede automaticamente nomeação e posse em concurso público'],
        ['Súmula Vinculante 18', 'Dissolução conjugal no curso do mandato não afasta inelegibilidade reflexa'],
    ])

    h(doc, '8 Dez questões de aplicação', 1)
    p(doc, 'Resolva antes de consultar o comentário. As quatro primeiras questões são fáceis, as quatro seguintes são médias e as duas últimas são difíceis.')
    for q in QUESTIONS:
        add_question(doc, q)

    h(doc, '9 Conteúdo resumido', 1)
    p(doc, 'Direitos sociais. O art. 6º reúne o núcleo geral, mas o tema também envolve proteção trabalhista, liberdade sindical, greve e representação. O art. 7º é um piso de direitos e contém números que precisam ser associados à regra correta.')
    p(doc, 'Nacionalidade. Brasileiro nato surge das hipóteses constitucionais de nascimento; naturalizado depende de procedimento derivado. As diferenças entre natos e naturalizados só podem vir da Constituição. A EC 131/2023 afastou a perda automática pela mera aquisição de outra nacionalidade.')
    p(doc, 'Direitos políticos. Capacidade ativa significa votar; passiva significa ser votado. Analfabeto vota facultativamente, mas é inelegível. A sequência de idades é 35, 30, 21 e 18. Reeleição ao mesmo cargo e candidatura a outro cargo seguem regras diferentes.')
    p(doc, 'Partidos. Primeiro vem a personalidade jurídica civil; depois o registro do estatuto no TSE. Coligações são vedadas nas eleições proporcionais. Cláusula de desempenho, fidelidade e percentuais de financiamento precisam ser estudados separadamente.')

    h(doc, '10 Cinco questões de revisão', 1)
    p(doc, 'Julgue certo ou errado e justifique em voz alta antes de consultar a resposta.')
    for q in REVIEW_QUESTIONS:
        add_question(doc, q, review=True)

    h(doc, '11 Revisão ativa', 1)
    h(doc, 'Revisão de cinco minutos', 2)
    bullets(doc, [
        'Recite os direitos do art. 6º sem consultar.',
        'Explique a transição da EC 141/2026 e diferencie regra permanente de aplicação temporal.',
        'Reconstrua as quatro hipóteses principais de brasileiro nato.',
        'Diga os sete cargos privativos de nato.',
        'Associe 35, 30, 21 e 18 aos cargos eletivos.',
        'Compare reeleição ao mesmo cargo com candidatura a outro cargo.',
        'Explique personalidade civil e registro do estatuto partidário no TSE.',
    ])

    h(doc, 'O módulo em um minuto', 2)
    p(doc, 'Direitos sociais concretizam igualdade material e incluem prestações, proteção ao trabalho e liberdades coletivas. Nacionalidade exige atenção ao local do nascimento, aos pais e às hipóteses constitucionais de registro, residência e opção. Direitos políticos separam votar de ser votado, com requisitos, idades e inelegibilidades. Partidos possuem autonomia, mas devem respeitar limites democráticos, registro, fidelidade, desempenho e regras de financiamento.')

    h(doc, 'O que decorar', 2)
    bullets(doc, [
        'Art. 11: mais de 200 empregados.',
        'Prescrição trabalhista: cinco anos, ação até dois anos após a extinção.',
        'Trabalho: regra aos 16; aprendiz aos 14; proteção especial até os 18.',
        'Naturalização extraordinária: mais de quinze anos, sem condenação penal e mediante requerimento.',
        'Idades eleitorais: 35, 30, 21 e 18.',
        'Renúncia para outro cargo: seis meses.',
        'AIME: quinze dias da diplomação.',
        'Consulta local: envio até noventa dias antes.',
        'Anterioridade eleitoral: um ano.',
        'Partidos: 5% para programas de participação feminina e mínimos constitucionais de 30% nas destinações de campanha indicadas pelo art. 17.',
    ])

    h(doc, 'Pegadinhas que a prova usa', 2)
    p(doc, '“Todos os direitos sociais são prestações do Estado.” Errado: liberdade sindical e greve são exemplos de liberdades coletivas.')
    p(doc, '“Empresa com 200 empregados deve eleger representante.” Errado: a Constituição exige mais de 200.')
    p(doc, '“A dupla nacionalidade extingue a brasileira.” Regra desatualizada após a EC 131/2023.')
    p(doc, '“Analfabeto não vota.” Errado: seu voto é facultativo; a restrição recai sobre a elegibilidade.')
    p(doc, '“Chefe do Executivo sempre renuncia para concorrer.” Errado: a renúncia de seis meses é exigida para outro cargo, não para a reeleição ao mesmo cargo.')
    p(doc, '“Partido nasce juridicamente com o registro no TSE.” Errado: primeiro adquire personalidade na forma da lei civil.')

    h(doc, 'Erros de estudo a corrigir', 2)
    p(doc, 'Decorar listas sem relacioná-las a um caso. Correção: para cada regra, invente uma situação em que uma palavra ou prazo mude o resultado.')
    p(doc, 'Misturar capacidade ativa e passiva. Correção: pergunte separadamente se a pessoa pode votar e se pode ser votada.')
    p(doc, 'Ignorar a data de atualização. Correção: em matéria constitucional recente, verifique a redação vigente e a norma de transição.')
    p(doc, 'Tratar jurisprudência como frase absoluta. Correção: registre o fato analisado, a tese e o limite de aplicação.')

    h(doc, 'Checklist de domínio', 2)
    for item in [
        '☐ Explico por que direitos sociais não se resumem a prestações estatais.',
        '☐ Reconstruo os números principais do art. 7º e a transição da EC 141/2026.',
        '☐ Distingo registro sindical de autorização estatal e contribuição de filiação.',
        '☐ Explico as quatro hipóteses principais de brasileiro nato.',
        '☐ Separo brasileiro nato naturalizado e português equiparado.',
        '☐ Identifico os cargos privativos de nato e as regras de extradição.',
        '☐ Diferencio sufrágio voto alistabilidade e elegibilidade.',
        '☐ Aplico as idades 35 30 21 e 18 sem trocar os cargos.',
        '☐ Distingo reeleição ao mesmo cargo de candidatura a outro cargo.',
        '☐ Explico personalidade registro desempenho fidelidade e financiamento partidário.',
        '☐ Justifico as respostas das quinze questões do módulo.',
    ]:
        p(doc, item)
    p(doc, 'Se algum item ainda depender de adivinhação, retorne à aula correspondente antes de marcar o módulo como concluído.')

    h(doc, '12 Referências e fontes de consulta', 1)
    p(doc, 'Constituição da República Federativa do Brasil de 1988. Texto compilado, especialmente arts. 6º a 17. Consulta em 22 de outubro de 2026.')
    p(doc, 'Fonte oficial: https://www.planalto.gov.br/ccivil_03/constituicao/constituicao.htm')
    p(doc, 'Emenda Constitucional nº 141, de 2026. Jornada de trabalho e repouso semanal.')
    p(doc, 'Fonte oficial: https://www.planalto.gov.br/ccivil_03/constituicao/emendas/emc/emc141.htm')
    p(doc, 'Emenda Constitucional nº 131, de 2023. Perda e reaquisição da nacionalidade.')
    p(doc, 'Fonte oficial: https://www.planalto.gov.br/ccivil_03/constituicao/emendas/emc/emc131.htm')
    p(doc, 'Emenda Constitucional nº 133, de 2024. Recursos partidários para candidaturas de pessoas pretas e pardas.')
    p(doc, 'Fonte oficial: https://www.planalto.gov.br/ccivil_03/constituicao/emendas/emc/emc133.htm')
    p(doc, 'Supremo Tribunal Federal. Temas 370, 541, 935 e 1.190 e Súmula Vinculante 18.')
    p(doc, 'Fonte oficial: https://portal.stf.jus.br/jurisprudenciaRepercussao/tema.asp?num=1190')
    p(doc, 'Base de partida: arquivo CF M03 enviado pelo usuário. O conteúdo foi reorganizado e ampliado pedagogicamente, com preservação dos temas e atualização jurídica indicada nas fontes oficiais.')


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    doc = Document(str(TEMPLATE))
    B.clear_body(doc)
    update_running_text(doc)
    add_cover(doc)
    build_body(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(json.dumps({'output': str(OUT), 'bytes': OUT.stat().st_size, 'questions': 15}, ensure_ascii=False))


if __name__ == '__main__':
    main()
