import json
import re
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path('/workspace/scratch/150b36589797/upload/01-CF-M02-Apostila-Didatica-1.docx')
SOURCE = Path('/workspace/scratch/150b36589797/output/cf-m01/CF-M01-Teoria-completa.txt')
DATA_JS = ROOT / 'content/cf/m01-native-data.js'
OUT = Path('/workspace/scratch/150b36589797/output/cf-m01/CF-M01-Apostila-Didatica-1.docx')

def load_native():
    text = DATA_JS.read_text(encoding='utf-8')
    match = re.search(r'cf\.m01=(\{.*\});\}\)\(window\);?\s*$', text, re.S)
    if not match:
        raise RuntimeError('Não foi possível ler os dados nativos do CF M01')
    return json.loads(match.group(1))

def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_borders(cell, color='D9D9D9', size='6'):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn('w:tcBorders'))
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), size)
        node.set(qn('w:space'), '0')
        node.set(qn('w:color'), color)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    cant = OxmlElement('w:cantSplit')
    trPr.append(cant)

def hyperlink(paragraph, text, url):
    part = paragraph.part
    rel = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), rel)
    run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '365F84'); rPr.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); rPr.append(underline)
    run.append(rPr)
    t = OxmlElement('w:t'); t.text = text; run.append(t); link.append(run)
    paragraph._p.append(link)

def add_run_format(run, size=11, bold=False, color=None):
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def add_rich_paragraph(doc, text, style='Normal'):
    text = text.strip()
    if not text:
        return None
    if text.startswith('Fonte oficial:') and 'https://' in text:
        url = text.split('https://', 1)[1]
        url = 'https://' + url.strip()
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(2)
        hyperlink(p, 'Abrir fonte oficial', url)
        return p
    p = doc.add_paragraph(style=style)
    p.paragraph_format.widow_control = True
    labels = ('Exemplo explicado.', 'Caso explicado.', 'Caso resolvido.', 'Caso de comparação.', 'Aprofundamento.', 'Teste de compreensão.', 'Resposta comentada.', 'Guarde o essencial.')
    lead = next((label for label in labels if text.startswith(label)), None)
    if lead:
        r = p.add_run(lead + ' ')
        add_run_format(r, bold=True)
        r = p.add_run(text[len(lead):].lstrip())
        add_run_format(r)
    else:
        r = p.add_run(text)
        add_run_format(r)
    return p

def cells_for(line):
    if not re.search(r'\S {3,}\S', line):
        return None
    cells = [c.strip() for c in re.split(r' {3,}', line.strip()) if c.strip()]
    return cells if len(cells) >= 2 else None

def add_table(doc, rows):
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = 'Normal Table'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for ri, row_data in enumerate(rows):
        row = table.rows[ri]
        no_split(row)
        for ci in range(width):
            cell = row.cells[ci]
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            value = row_data[ci] if ci < len(row_data) else ''
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = 0
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            add_run_format(run, size=9.5, bold=(ri == 0), color='FFFFFF' if ri == 0 else None)
            set_cell_shading(cell, '142F50' if ri == 0 else ('F2F5F7' if ri % 2 == 0 else 'FFFFFF'))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table

def add_blocks(doc, text):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        if not lines[i].strip():
            i += 1; continue
        line = lines[i].strip()
        # A block of aligned rows becomes a readable comparison table.
        block = []
        j = i
        while j < len(lines) and lines[j].strip():
            block.append(lines[j].strip()); j += 1
        table_rows = [cells_for(x) for x in block]
        if len(block) >= 2 and all(row is not None for row in table_rows):
            add_table(doc, table_rows)
            i = j; continue
        add_rich_paragraph(doc, line, 'List Bullet' if line.startswith('•') else 'Normal')
        i += 1

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style='Heading 1' if level == 1 else 'Heading 2')
    p.paragraph_format.keep_with_next = True
    p.add_run(text.strip())
    return p

def add_question(doc, q, review=False):
    prefix = 'Revisão ' if review else 'Questão '
    add_heading(doc, f"{prefix}{q['n']} {q['difficulty']}", 2)
    add_rich_paragraph(doc, q['q'])
    for idx, option in enumerate(q['options']):
        add_rich_paragraph(doc, f"{chr(65 + idx)}) {option}")
    add_rich_paragraph(doc, f"Gabarito comentado: {chr(65 + q['answer'])}. {q['explanation']}")

def add_title(doc, data):
    p = doc.add_paragraph(style='Normal'); r = p.add_run('BASE COMPLETA'); add_run_format(r, size=12, bold=True, color='142F50')
    p.paragraph_format.space_after = Pt(18)
    doc.add_paragraph(data['title'], style='Title')
    doc.add_paragraph('Direito Constitucional • CF M01', style='Subtitle')
    doc.add_paragraph('Teoria explicada e aprofundamento para concursos de tribunais', style='Subtitle')
    p = doc.add_paragraph('Edição didática • 21 de outubro de 2026', style='Normal')
    p.paragraph_format.space_before = Pt(10)
    p = doc.add_paragraph('Aprenda a compreender como a Constituição organiza o poder, produz efeitos e orienta a solução de casos jurídicos.', style='Normal')
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(20)

def split_summary(summary):
    review = summary.find('\nREVISÃO ATIVA')
    refs = summary.find('\nREFERÊNCIAS E FONTES DE CONSULTA')
    checklist = summary.find('\nCHECKLIST DE DOMÍNIO')
    return summary[:review].strip(), summary[review + 1:refs].strip(), summary[refs + 1:checklist].strip(), summary[checklist + 1:].strip()

def main():
    data = load_native()
    complete = SOURCE.read_text(encoding='utf-8')
    # The full text already contains theory and the law-reading section. The active review and checklist are added below from the native data.
    review_marker = complete.find('\n12. REVISÃO ATIVA')
    theory = complete[:review_marker].strip()
    summary_main, review_text, references, checklist = split_summary(data['summary'])

    doc = Document(str(TEMPLATE))
    clear_body(doc)

    # Keep the template's running header/footer styling, but identify this
    # module correctly.  The page-number field in the footer remains intact.
    for section in doc.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                for run in paragraph.runs:
                    run.text = run.text.replace('CF M02', 'CF M01')

    add_title(doc, data)

    # Intro and theory up to the law-reading section.
    blocks = theory.split('\n\n')
    for block in blocks:
        b = block.strip()
        if not b: continue
        if b == 'LEITURA ORIENTADA': add_heading(doc, 'Leitura orientada', 1); continue
        if b == 'PERCURSO DO MÓDULO': add_heading(doc, 'Percurso do módulo', 2); continue
        if b == 'CHECKLIST DE DOMÍNIO': continue
        if re.match(r'^\d+\. [A-ZÁÉÍÓÚÂÊÔÃÕÇ0-9]', b):
            add_heading(doc, b.title(), 1); continue
        if re.match(r'^\d+\.\d+\. ', b):
            add_heading(doc, b.title(), 2); continue
        add_blocks(doc, b)

    add_heading(doc, '12 Dez questões de aplicação', 1)
    add_rich_paragraph(doc, 'Resolva antes de consultar o gabarito. Há quatro questões fáceis, quatro médias e duas difíceis. A dificuldade é uma classificação didática desta edição.')
    for q in data['questions']['complete']:
        add_question(doc, q)

    add_heading(doc, '13 Conteúdo resumido', 1)
    add_blocks(doc, summary_main.replace('CONTEÚDO RESUMIDO — PRIMEIRA LEITURA E REVISÃO', '').strip())

    add_heading(doc, '14 Cinco questões de revisão', 1)
    add_rich_paragraph(doc, 'Julgue certo ou errado. Há duas questões fáceis, duas médias e uma difícil. Explique o motivo antes de consultar as respostas.')
    for q in data['questions']['summary']:
        add_question(doc, q, review=True)

    add_heading(doc, '15 Revisão ativa', 1)
    add_blocks(doc, review_text.replace('REVISÃO ATIVA', '').strip())

    add_heading(doc, '16 Referências e fontes de consulta', 1)
    add_blocks(doc, references.replace('REFERÊNCIAS E FONTES DE CONSULTA', '').strip())

    add_heading(doc, 'Checklist de domínio', 2)
    for line in checklist.splitlines():
        if line.strip():
            text = line.strip()
            if text.startswith('• ☐'): text = '☐' + text[3:]
            # The reference M02 uses plain paragraphs for the checklist,
            # leaving the checkbox glyph itself as the only marker.
            add_rich_paragraph(doc, text, 'Normal')

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(json.dumps({'output': str(OUT), 'bytes': OUT.stat().st_size, 'questions': 15}, ensure_ascii=False))

if __name__ == '__main__':
    main()
