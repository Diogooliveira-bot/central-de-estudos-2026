"""M06: fonte editorial única para a apostila e o conteúdo nativo da Central."""
import argparse
import importlib.util
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
WORKSPACE = ROOT.parent
EDITORIAL = ROOT / 'content/cf/editorial'
OUT = WORKSPACE / 'output/cf-m06'


def base_builder():
    spec = importlib.util.spec_from_file_location('cf_docx_base', ROOT / 'scripts/build-cf-m01-docx.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def questions_text(questions, review=False):
    blocks = []
    for q in questions:
        prefix = 'REVISÃO' if review else 'QUESTÃO'
        blocks.append(f"{prefix} {q['n']} {q['difficulty'].upper()}")
        blocks.append(q['q'])
        blocks.extend(f"{chr(65 + i)}) {v}" for i, v in enumerate(q['options']))
        blocks.append(f"Resposta comentada. Gabarito {chr(65 + q['answer'])}. {q['explanation']}")
    return '\n\n'.join(blocks)


def text_content():
    questions = json.loads((EDITORIAL / 'm06-questions.json').read_text())
    text = (EDITORIAL / 'm06-apostila.txt').read_text()
    text = text.replace('{{QUESTIONS}}', questions_text(questions['complete']))
    text = text.replace('{{REVIEW_QUESTIONS}}', questions_text(questions['summary'], True))
    assert '{{' not in text
    return text, questions


def build_docx(text, template):
    b = base_builder()
    doc = Document(template)
    b.clear_body(doc)
    for section in doc.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                for run in paragraph.runs:
                    run.text = run.text.replace('CF M02', 'CF M06')
    for style_name in ('Title', 'Subtitle', 'Heading 1', 'Heading 2'):
        doc.styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    doc.core_properties.title = 'Organização dos Poderes e Poder Legislativo'
    doc.core_properties.subject = 'CF M06 Apostila didática'
    doc.core_properties.author = 'Base Completa'
    doc.core_properties.keywords = 'Direito Constitucional, Congresso Nacional, CPI, imunidades, CF M06'
    p = doc.add_paragraph('BASE COMPLETA')
    p.paragraph_format.space_after = Pt(18)
    b.add_run_format(p.runs[0], size=12, bold=True, color='142F50')
    doc.add_paragraph('Organização dos Poderes e Poder Legislativo', style='Title')
    doc.add_paragraph('Direito Constitucional • CF M06', style='Subtitle')
    doc.add_paragraph('Teoria explicada e aprofundamento para concursos de tribunais', style='Subtitle')
    b.add_rich_paragraph(doc, 'Edição didática • 22 de outubro de 2026')
    b.add_rich_paragraph(doc, 'Aprenda a identificar quem decide cada matéria, como o Congresso funciona e quais garantias e responsabilidades acompanham o mandato parlamentar.')

    for block in text.split('\n\n'):
        lines = [line.strip() for line in block.strip().splitlines() if line.strip()]
        if not lines:
            continue
        rows = [b.cells_for(line) for line in lines]
        if len(rows) > 1 and all(rows):
            assert len(set(map(len, rows))) == 1, rows
            table = b.add_table(doc, rows)
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        else:
            for line in lines:
                if re.match(r'^\d+(?:\.\d+)?\s', line):
                    b.add_heading(doc, line, 2 if re.match(r'^\d+\.', line) else 1)
                elif line.isupper() and len(line) < 100 and not line.startswith('•'):
                    b.add_heading(doc, line, 2)
                else:
                    # The interactive boxes are rendered by the Central reader.
                    b.add_rich_paragraph(doc, line.removeprefix('• '), 'List Bullet' if line.startswith('• ') and '☐' not in line else 'Normal')
    # Keep each exercise together, including its explanation, in the printed edition.
    in_question = False
    for paragraph in doc.paragraphs:
        if re.match(r'^(QUESTÃO|REVISÃO) \d+ ', paragraph.text):
            in_question = True
        if in_question:
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = not paragraph.text.startswith('Resposta comentada.')
            if paragraph.text.startswith('Resposta comentada.'):
                in_question = False
        if paragraph.text == 'CHECKLIST DE DOMÍNIO':
            paragraph.paragraph_format.page_break_before = True
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / 'CF-M06-Apostila-Didatica.docx'
    doc.save(path)
    return path


def integrate(text, questions):
    path = ROOT / 'content/cf/m06-native-data.js'
    source = path.read_text()
    match = re.search(r'\.m06=(\{.*\});\}\)\(window\);\s*$', source, re.S)
    data = json.loads(match.group(1))
    data['complete'] = text
    data['questions'] = questions
    data['lawTopicId'] = 'cf-m06'
    data['editorialVersion'] = '20261022cfm06apostila1'
    payload = json.dumps(data, ensure_ascii=False, separators=(',', ':'))
    path.write_text("(function(g){'use strict';g.BASE_NATIVE_CONTENT=g.BASE_NATIVE_CONTENT||{};g.BASE_NATIVE_CONTENT.cf=g.BASE_NATIVE_CONTENT.cf||{};g.BASE_NATIVE_CONTENT.cf.m06=" + payload + ';})(window);\n')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--template', type=Path, default=WORKSPACE / 'upload/01-CF-M02-Apostila-Didatica-1.docx')
    parser.add_argument('--integrate', action='store_true')
    args = parser.parse_args()
    content, questions = text_content()
    output = build_docx(content, args.template)
    if args.integrate:
        integrate(content, questions)
    (OUT / 'CF-M06-Teoria-completa.txt').write_text(content)
    print(json.dumps({'docx': str(output), 'characters': len(content), 'questions': 15}, ensure_ascii=False))
