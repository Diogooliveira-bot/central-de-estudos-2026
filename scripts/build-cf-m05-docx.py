import importlib.util
import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path('/workspace/scratch/150b36589797/upload/01-CF-M02-Apostila-Didatica-1.docx')
SOURCE = Path('/workspace/scratch/150b36589797/upload/CF-M05-Administracao-Publica-e-servidores-publicos-Teoria-completa.txt')
OUT = Path('/workspace/scratch/150b36589797/output/cf-m05/CF-M05-Apostila-Didatica-1.docx')


def load_base_builder():
    path = ROOT / 'scripts/build-cf-m01-docx.py'
    spec = importlib.util.spec_from_file_location('cf_m01_docx_builder', path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


B = load_base_builder()


def p(doc, text, style='Normal'):
    return B.add_rich_paragraph(doc, text, style)


def h(doc, text, level=1):
    return B.add_heading(doc, text, level)


def bullets(doc, items):
    for item in items:
        p(doc, item, 'List Bullet')


def table(doc, headers, rows):
    tbl = B.add_table(doc, [headers, *rows])
    header_props = tbl.rows[0]._tr.get_or_add_trPr()
    header_props.append(OxmlElement('w:tblHeader'))
    for row in tbl.rows:
        row_props = row._tr.get_or_add_trPr()
        row_props.append(OxmlElement('w:cantSplit'))
    return tbl


def update_running_text(doc):
    for section in doc.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                for run in paragraph.runs:
                    run.text = run.text.replace('CF M02', 'CF M05')


def add_cover(doc):
    top = doc.add_paragraph(style='Normal')
    run = top.add_run('BASE COMPLETA')
    B.add_run_format(run, size=12, bold=True, color='142F50')
    top.paragraph_format.space_after = Pt(18)
    doc.add_paragraph('Administração Pública e servidores públicos', style='Title')
    doc.add_paragraph('Direito Constitucional • CF M05', style='Subtitle')
    doc.add_paragraph('Teoria explicada e aprofundamento para concursos de tribunais', style='Subtitle')
    meta = doc.add_paragraph('Edição didática • 22 de outubro de 2026', style='Normal')
    meta.paragraph_format.space_before = Pt(10)
    intro = doc.add_paragraph(
        'Entenda os princípios do art. 37, o concurso público, os vínculos funcionais, a responsabilidade estatal, a estabilidade e a previdência dos servidores.',
        style='Normal',
    )
    intro.paragraph_format.space_before = Pt(18)
    intro.paragraph_format.space_after = Pt(20)


QUESTIONS = [
    {
        'n': 1, 'difficulty': 'Fácil',
        'q': 'Qual alternativa traduz corretamente a legalidade administrativa?',
        'options': [
            'O agente público pode fazer tudo o que a lei não proíbe.',
            'O agente público atua com fundamento e dentro dos limites do ordenamento jurídico.',
            'A legalidade permite afastar a lei quando o resultado parecer eficiente.',
            'A legalidade alcança apenas a Administração direta federal.',
        ],
        'answer': 1,
        'explanation': 'Para a Administração, liberdade de agir não é a regra. A atuação precisa encontrar fundamento jurídico e respeitar finalidade, competência e procedimento.'
    },
    {
        'n': 2, 'difficulty': 'Fácil',
        'q': 'A função de confiança pode ser exercida por quem?',
        'options': [
            'Por qualquer pessoa escolhida livremente.',
            'Exclusivamente por servidor ocupante de cargo efetivo.',
            'Somente por empregado de empresa pública.',
            'Exclusivamente por agente político eleito.',
        ],
        'answer': 1,
        'explanation': 'O art. 37, V reserva a função de confiança ao servidor efetivo. Cargo em comissão admite livre nomeação, observadas a lei e as atribuições de direção, chefia e assessoramento.'
    },
    {
        'n': 3, 'difficulty': 'Fácil',
        'q': 'Qual é o prazo constitucional de validade de um concurso público?',
        'options': [
            'Exatamente dois anos, sem prorrogação.',
            'Até dois anos, prorrogável uma vez por igual período.',
            'Até quatro anos, prorrogável livremente.',
            'Prazo indeterminado até o preenchimento das vagas.',
        ],
        'answer': 1,
        'explanation': 'A Constituição fixa o máximo de dois anos e admite uma única prorrogação por período igual ao inicialmente definido.'
    },
    {
        'n': 4, 'difficulty': 'Fácil',
        'q': 'Quem adquire a estabilidade prevista no art. 41?',
        'options': [
            'Todo agente público após três anos.',
            'O empregado público aprovado em concurso.',
            'O servidor nomeado para cargo efetivo por concurso, após três anos e avaliação especial.',
            'O ocupante de cargo em comissão depois do estágio probatório.',
        ],
        'answer': 2,
        'explanation': 'A estabilidade constitucional exige cargo efetivo, concurso, três anos de efetivo exercício e avaliação especial de desempenho.'
    },
    {
        'n': 5, 'difficulty': 'Média',
        'q': 'Candidato aprovado dentro das vagas deixa de ser nomeado porque os cargos foram extintos após superação do limite prudencial de gastos. Qual afirmação segue o Tema 1.164 do STF?',
        'options': [
            'O direito à nomeação jamais pode ser mitigado.',
            'A extinção sempre afasta o direito, mesmo sem motivação.',
            'A mitigação pode ocorrer excepcionalmente se a extinção for anterior ao fim da validade e devidamente motivada.',
            'Basta uma decisão administrativa tomada depois do vencimento do concurso.',
        ],
        'answer': 2,
        'explanation': 'A regra é o direito subjetivo dentro das vagas. O Tema 1.164 reconheceu exceção estrita, ligada ao limite prudencial, à extinção superveniente, à anterioridade e à motivação adequada.'
    },
    {
        'n': 6, 'difficulty': 'Média',
        'q': 'Após a EC 138/2025, qual acumulação é admitida, desde que haja compatibilidade de horários?',
        'options': [
            'Um cargo de professor com outro de qualquer natureza.',
            'Três cargos de professor.',
            'Dois cargos administrativos de qualquer natureza.',
            'Um cargo em comissão com dois cargos efetivos.',
        ],
        'answer': 0,
        'explanation': 'A nova redação do art. 37, XVI, b permite um cargo de professor com outro de qualquer natureza. A compatibilidade de horários continua indispensável.'
    },
    {
        'n': 7, 'difficulty': 'Média',
        'q': 'Durante greve de servidores, qual é a regra do Tema 531 do STF?',
        'options': [
            'Os dias parados nunca podem ser descontados.',
            'A Administração deve descontar os dias, admitida compensação; não há desconto se a greve decorreu de conduta ilícita do Poder Público.',
            'O desconto depende sempre de sentença judicial prévia.',
            'A remuneração deve ser dobrada se houver acordo posterior.',
        ],
        'answer': 1,
        'explanation': 'A paralisação suspende o vínculo funcional para esse efeito. A tese admite compensação por acordo e afasta o desconto quando demonstrada conduta ilícita do Poder Público.'
    },
    {
        'n': 8, 'difficulty': 'Média',
        'q': 'Um agente público causa dano a terceiro no exercício de suas funções. Em regra, como se organiza a responsabilidade?',
        'options': [
            'A vítima precisa provar culpa do agente para demandar o Estado.',
            'O Estado responde objetivamente e o regresso contra o agente exige dolo ou culpa.',
            'Somente o agente responde, sempre de forma objetiva.',
            'O Estado responde apenas se houver condenação penal do agente.',
        ],
        'answer': 1,
        'explanation': 'Na relação vítima e Estado ou prestador de serviço público, aplica-se a responsabilidade objetiva do art. 37, § 6º. A ação regressiva contra o agente depende de dolo ou culpa.'
    },
    {
        'n': 9, 'difficulty': 'Difícil',
        'q': 'Servidor estável é demitido, mas a demissão é invalidada por sentença. Outro servidor estável ocupa a vaga. Qual solução constitucional é correta?',
        'options': [
            'O demitido retorna por reversão e o ocupante recebe indenização.',
            'O demitido é reintegrado; o ocupante estável pode ser reconduzido, aproveitado ou posto em disponibilidade, sem indenização.',
            'Ambos permanecem no mesmo cargo até surgir nova vaga.',
            'O ocupante da vaga perde automaticamente qualquer vínculo com a Administração.',
        ],
        'answer': 1,
        'explanation': 'O art. 41, § 2º disciplina reintegração do demitido e as três soluções para o ocupante estável: recondução, aproveitamento ou disponibilidade proporcional, sem indenização.'
    },
    {
        'n': 10, 'difficulty': 'Difícil',
        'q': 'Qual afirmação está atualizada sobre o art. 39 após o julgamento definitivo da ADI 2.135?',
        'options': [
            'A Constituição exige obrigatoriamente um único regime jurídico em todos os entes.',
            'O STF declarou inconstitucional toda a EC 19/1998.',
            'O STF reconheceu a validade da alteração que retirou a obrigatoriedade constitucional do regime jurídico único, preservadas as situações consolidadas.',
            'A decisão transformou automaticamente todos os servidores estatutários em empregados celetistas.',
        ],
        'answer': 2,
        'explanation': 'Em novembro de 2024, o STF julgou válida a alteração promovida pela EC 19/1998. A decisão não converte vínculos existentes automaticamente nem dispensa a disciplina legal de cada ente.'
    },
]


REVIEW_QUESTIONS = [
    {
        'n': 1, 'difficulty': 'Fácil',
        'q': 'Certo ou errado: publicidade institucional pode usar nome e imagem da autoridade se a obra tiver interesse coletivo.',
        'options': ['Certo.', 'Errado.'], 'answer': 1,
        'explanation': 'A publicidade deve ter caráter educativo, informativo ou de orientação social e não pode caracterizar promoção pessoal.'
    },
    {
        'n': 2, 'difficulty': 'Fácil',
        'q': 'Certo ou errado: contratação temporária exige lei e necessidade temporária de excepcional interesse público.',
        'options': ['Certo.', 'Errado.'], 'answer': 0,
        'explanation': 'Esses elementos limitam a exceção do art. 37, IX. Necessidades permanentes não podem ser cobertas indefinidamente por contratos precários.'
    },
    {
        'n': 3, 'difficulty': 'Média',
        'q': 'Certo ou errado: o teto remuneratório incide sempre sobre a soma dos vínculos acumulados licitamente.',
        'options': ['Certo.', 'Errado.'], 'answer': 1,
        'explanation': 'O STF entende que, nas acumulações constitucionalmente autorizadas, o teto é verificado em regra em cada vínculo separadamente.'
    },
    {
        'n': 4, 'difficulty': 'Média',
        'q': 'Certo ou errado: improbidade administrativa gera perda dos direitos políticos.',
        'options': ['Certo.', 'Errado.'], 'answer': 1,
        'explanation': 'O art. 37, § 4º prevê suspensão dos direitos políticos, não perda. A tipificação atual dos arts. 9º, 10 e 11 da LIA exige dolo.'
    },
    {
        'n': 5, 'difficulty': 'Difícil',
        'q': 'Certo ou errado: todo ocupante de cargo em comissão integra obrigatoriamente um RPPS.',
        'options': ['Certo.', 'Errado.'], 'answer': 1,
        'explanation': 'O ocupante exclusivamente de cargo em comissão, assim como empregado público e temporário, submete-se em regra ao RGPS. O RPPS alcança titular de cargo efetivo.'
    },
]


def add_question(doc, q, review=False):
    if not review and q['n'] == 8:
        doc.add_page_break()
    label = 'Revisão' if review else 'Questão'
    h(doc, f"{label} {q['n']} {q['difficulty']}", 2)
    p(doc, q['q'])
    for i, option in enumerate(q['options']):
        p(doc, f"{chr(65 + i)}) {option}")
    p(doc, f"Resposta comentada. {chr(65 + q['answer'])}. {q['explanation']}")


def build_body(doc):
    h(doc, 'Leitura orientada', 1)
    p(doc, 'Imagine três situações. Um candidato aprovado dentro das vagas aguarda nomeação. Um gestor deseja contratar pessoas sem concurso para uma necessidade permanente. Um servidor causa dano a um cidadão durante o serviço. As respostas não vêm de uma única regra: é preciso identificar o princípio, o tipo de vínculo, a exceção constitucional e, quando houver, a tese do STF.')
    p(doc, 'O módulo foi organizado como uma sequência de decisões. Primeiro você compreende quem compõe a Administração e quais princípios controlam sua atuação. Depois classifica cargos, empregos e funções, analisa ingresso, remuneração e acumulação. Ao final, conecta responsabilidade, improbidade, estabilidade e previdência.')
    p(doc, 'Se estiver começando, use os exemplos e as regras de bolso. Se já possui base, concentre-se nas atualizações de 2024 a 2026: ADI 2.135, Tema 1.164, EC 138/2025 e EC 140/2026.')

    h(doc, 'Percurso do módulo', 2)
    bullets(doc, [
        'Aula 1 • Administração em sentido subjetivo e objetivo.',
        'Aula 2 • Princípios expressos e implícitos.',
        'Aula 3 • Cargos empregos funções e concurso público.',
        'Aula 4 • Confiança comissão temporários e nepotismo.',
        'Aula 5 • Greve remuneração teto e acumulação.',
        'Aula 6 • Mandato eletivo e regime constitucional.',
        'Aula 7 • Responsabilidade civil e improbidade.',
        'Aula 8 • Estabilidade e movimentações funcionais.',
        'Aula 9 • RPPS e atualizações constitucionais.',
        'Aula 10 • Questões revisão ativa e checklist.',
    ])

    h(doc, '1 Como enxergar a Administração Pública', 1)
    h(doc, '1.1 Quem atua e o que é feito', 2)
    p(doc, 'A função administrativa é típica do Executivo, mas Legislativo e Judiciário também a exercem quando organizam pessoal, contratos e serviços internos. Por isso, o art. 37 alcança a Administração direta e indireta de todos os Poderes e de todos os entes federativos.')
    table(doc, ['Sentido', 'Pergunta', 'Conteúdo'], [
        ['Subjetivo ou orgânico', 'Quem exerce a atividade', 'Órgãos e entidades administrativas'],
        ['Objetivo ou material', 'O que a Administração faz', 'Serviço público polícia administrativa fomento e intervenção'],
    ])
    p(doc, 'Caso explicado. Um tribunal realiza concurso para servidores. Julgar processos é sua função típica jurisdicional; organizar o concurso e gerir pessoal é atividade administrativa atípica, sujeita ao art. 37.')

    h(doc, '1.2 Administração direta e indireta', 2)
    table(doc, ['Estrutura', 'Composição', 'Personalidade jurídica'], [
        ['Administração direta', 'Órgãos da União Estados DF e Municípios', 'O órgão não possui personalidade própria'],
        ['Administração indireta', 'Autarquias fundações empresas públicas e sociedades de economia mista', 'A entidade possui personalidade própria'],
    ])
    p(doc, 'A autarquia é criada diretamente por lei específica. Empresa pública, sociedade de economia mista e fundação têm instituição autorizada por lei específica; lei complementar define as áreas de atuação das fundações. Não troque criação com autorização.')

    doc.add_page_break()
    h(doc, '2 Princípios constitucionais', 1)
    h(doc, '2.1 O LIMPE explicado', 2)
    table(doc, ['Princípio', 'O que exige', 'Erro comum'], [
        ['Legalidade', 'Atuação com fundamento e limites jurídicos', 'Aplicar ao gestor a liberdade ampla do particular'],
        ['Impessoalidade', 'Finalidade pública isonomia e neutralidade', 'Usar publicidade oficial para promoção pessoal'],
        ['Moralidade', 'Probidade boa fé e ética juridicamente exigível', 'Reduzir a moralidade a opinião subjetiva'],
        ['Publicidade', 'Transparência divulgação e controle', 'Tratá-la como absoluta mesmo diante de sigilo legítimo'],
        ['Eficiência', 'Qualidade resultado e uso racional de recursos', 'Invocá-la para ignorar legalidade e procedimento'],
    ])
    p(doc, 'Os princípios não funcionam isoladamente. Um ato eficiente, mas praticado sem competência ou para favorecer alguém, continua inválido. A publicidade institucional deve informar, educar ou orientar e não pode promover autoridades ou servidores.')
    p(doc, 'Entre os princípios implícitos mais cobrados estão razoabilidade, proporcionalidade, segurança jurídica, autotutela, continuidade e indisponibilidade do interesse público. Eles ajudam a interpretar o sistema, mas não substituem a literalidade constitucional.')

    doc.add_page_break()
    h(doc, '3 Pessoas vínculos e ingresso', 1)
    h(doc, '3.1 Cargo emprego e função', 2)
    table(doc, ['Vínculo', 'Regime básico', 'Ponto decisivo'], [
        ['Cargo efetivo', 'Estatutário', 'Ingresso por concurso e possibilidade de estabilidade'],
        ['Emprego público', 'Contratual celetista', 'Ingresso por concurso sem estabilidade do art. 41'],
        ['Função de confiança', 'Atribuição adicional', 'Exclusiva de servidor efetivo'],
        ['Cargo em comissão', 'Livre nomeação e exoneração', 'Somente direção chefia e assessoramento'],
        ['Contratação temporária', 'Vínculo excepcional previsto em lei', 'Necessidade temporária de excepcional interesse público'],
    ])
    p(doc, 'Requisitos de acesso precisam estar previstos em lei e guardar relação com o cargo. O edital detalha a lei, mas não pode criar restrição autônoma. Exame psicotécnico exige previsão legal; limite etário depende da natureza das atribuições; tatuagem só admite restrição em situações excepcionalíssimas ligadas ao conteúdo incompatível com valores constitucionais.')

    h(doc, '3.2 Concurso público', 2)
    p(doc, 'Cargo ou emprego público exige concurso de provas ou de provas e títulos, conforme a natureza e a complexidade. A exceção textual é o cargo em comissão declarado em lei de livre nomeação e exoneração. Formas de provimento que permitem ingresso em carreira diferente sem novo concurso são inconstitucionais.')
    table(doc, ['Pergunta de prova', 'Regra'], [
        ['Qual o prazo', 'Até dois anos'],
        ['Pode prorrogar', 'Uma vez por período igual'],
        ['Quem tem prioridade', 'Aprovado anterior durante o prazo previsto no edital'],
        ['Dentro das vagas', 'Direito subjetivo à nomeação como regra'],
        ['Cadastro de reserva', 'Expectativa em regra podendo virar direito diante de preterição arbitrária'],
    ])
    p(doc, 'Aprofundamento. O Tema 1.164 do STF admite mitigação excepcional do direito à nomeação dentro das vagas quando os cargos forem extintos por superação do limite prudencial de gastos com pessoal, antes do fim da validade e mediante motivação adequada. A exceção não autoriza conveniência genérica nem omissão administrativa.')

    h(doc, '3.3 Confiança comissão e nepotismo', 2)
    table(doc, ['Instituto', 'Quem pode exercer', 'Atribuições'], [
        ['Função de confiança', 'Exclusivamente servidor efetivo', 'Direção chefia e assessoramento'],
        ['Cargo em comissão', 'Livre nomeação com percentuais legais reservados à carreira', 'Direção chefia e assessoramento'],
    ])
    p(doc, 'Atividades técnicas, burocráticas ou operacionais permanentes não justificam cargo em comissão. Livre nomeação não significa liberdade para violar moralidade e impessoalidade. A Súmula Vinculante 13 combate nomeações de cônjuge, companheiro e parentes até o terceiro grau, inclusive por ajuste recíproco. Em cargos políticos, evite respostas absolutas: analise o caso, a qualificação, a subordinação e os princípios constitucionais.')

    h(doc, '3.4 Contratação temporária', 2)
    p(doc, 'O art. 37, IX permite que a lei estabeleça contratações por tempo determinado para necessidade temporária de excepcional interesse público. A sequência correta é lei autorizadora, necessidade transitória real, interesse público excepcional, prazo e adequação. Renovar indefinidamente contratos para atividade permanente burla o concurso.')

    h(doc, '4 Direitos funcionais remuneração e acumulação', 1)
    h(doc, '4.1 Associação e greve', 2)
    p(doc, 'O servidor civil possui direito à associação sindical. A greve deve observar os termos e limites de lei específica. Diante da omissão legislativa, o STF construiu disciplina para viabilizar o direito sem torná-lo absoluto. Pelo Tema 531, a Administração deve em regra descontar os dias parados, pode admitir compensação por acordo e não desconta quando a paralisação tiver sido provocada por conduta ilícita do Poder Público.')
    p(doc, 'Carreiras de segurança pública recebem tratamento mais restritivo em razão da continuidade e essencialidade das funções. Em prova, não aplique automaticamente a mesma solução a todo agente público.')

    h(doc, '4.2 Remuneração subsídio e teto', 2)
    p(doc, 'Remuneração e subsídio dependem de lei específica. Revisão geral anual também depende de lei e não cria, por si só, direito automático a reajuste em percentual determinado. Subsídio é pago em parcela única às categorias constitucionais, mas isso não elimina verbas genuinamente indenizatórias ou parcelas autônomas admitidas pela Constituição.')
    table(doc, ['Âmbito', 'Limite principal'], [
        ['União', 'Subsídio dos Ministros do STF'],
        ['Municípios', 'Subsídio do Prefeito'],
        ['Estados e DF no Executivo', 'Subsídio do Governador'],
        ['Estados e DF no Legislativo', 'Subsídio dos Deputados Estaduais ou Distritais'],
        ['Estados e DF no Judiciário', 'Subsídio dos Desembargadores do TJ limitado a 90,25 por cento do STF'],
    ])
    bullets(doc, [
        'Vinculação e equiparação remuneratória são vedadas, salvo autorização constitucional.',
        'Acréscimos não podem formar efeito cascata para novos acréscimos.',
        'Irredutibilidade convive com teto e incidência tributária.',
        'Verba chamada de indenizatória só fica fora do teto se possuir natureza realmente indenizatória.',
        'O teto alcança estatais dependentes de recursos públicos para pessoal ou custeio geral.',
    ])

    h(doc, '4.3 Acumulação remunerada atualizada', 2)
    p(doc, 'A regra continua sendo a vedação. A exceção exige duas etapas: enquadramento em hipótese constitucional e compatibilidade de horários. A EC 138/2025 alterou a segunda hipótese e permite um cargo de professor com outro de qualquer natureza.')
    table(doc, ['Hipótese permitida', 'Condição comum'], [
        ['Dois cargos de professor', 'Compatibilidade de horários'],
        ['Um cargo de professor com outro de qualquer natureza', 'Compatibilidade de horários'],
        ['Dois cargos ou empregos privativos de profissionais de saúde com profissões regulamentadas', 'Compatibilidade de horários'],
    ])
    p(doc, 'Nas acumulações constitucionalmente permitidas, o STF aplica em regra o teto a cada vínculo separadamente. A vedação de acumular abrange empregos e funções e alcança a Administração indireta e entidades controladas pelo Poder Público.')

    h(doc, '4.4 Servidor em mandato eletivo', 2)
    table(doc, ['Mandato', 'Situação funcional'], [
        ['Federal estadual ou distrital', 'Afastamento do cargo emprego ou função'],
        ['Prefeito', 'Afastamento com opção pela remuneração'],
        ['Vereador com horários compatíveis', 'Vantagens do vínculo e remuneração do mandato'],
        ['Vereador sem compatibilidade', 'Afastamento com direito de opção'],
    ])
    p(doc, 'Nos afastamentos previstos, o tempo conta para todos os efeitos legais, exceto promoção por merecimento. O servidor vinculado a RPPS permanece filiado ao regime de origem conforme a Constituição.')

    h(doc, '5 Regime constitucional e responsabilidade', 1)
    h(doc, '5.1 Artigo 39 após a ADI 2.135', 2)
    p(doc, 'Atenção à atualização. Em novembro de 2024, o STF concluiu o julgamento da ADI 2.135 e reconheceu a validade formal da alteração feita pela EC 19/1998, que retirou a obrigatoriedade constitucional do regime jurídico único. Isso supera a leitura baseada apenas na antiga cautelar. Cada ente precisa disciplinar seus vínculos por lei, e a decisão não converte automaticamente regimes já existentes.')
    p(doc, 'O art. 39 também organiza critérios remuneratórios, escolas de governo, direitos sociais aplicáveis aos ocupantes de cargo público e remuneração por subsídio. O Judiciário não pode aumentar vencimentos com fundamento apenas em isonomia, conforme a Súmula Vinculante 37.')

    h(doc, '5.2 Responsabilidade civil do Estado', 2)
    p(doc, 'Quando agente público, nessa qualidade, causa dano a terceiro, a vítima demanda a pessoa jurídica de direito público ou a prestadora privada de serviço público e não precisa provar culpa do agente. Deve demonstrar conduta, dano e nexo causal. O regime é o risco administrativo, que admite exame de excludentes e atenuantes.')
    table(doc, ['Relação', 'Regra de responsabilidade'], [
        ['Vítima contra Estado ou prestador', 'Objetiva nos termos do art. 37 § 6º'],
        ['Estado contra agente', 'Regressiva e dependente de dolo ou culpa'],
        ['Omissão estatal', 'Análise do dever específico de agir e das circunstâncias do caso'],
    ])
    p(doc, 'Exemplo. Viatura oficial conduzida em serviço causa acidente. A vítima prova dano, atuação estatal e nexo. Depois de indenizar, o Estado poderá cobrar do agente se demonstrar dolo ou culpa.')

    h(doc, '5.3 Improbidade administrativa', 2)
    p(doc, 'O art. 37, § 4º prevê suspensão dos direitos políticos, perda da função pública, indisponibilidade de bens e ressarcimento ao erário, na forma e gradação legais, sem prejuízo da ação penal cabível. A palavra correta é suspensão, não perda dos direitos políticos.')
    p(doc, 'No Tema 1.199, o STF exigiu dolo para tipificar os atos dos arts. 9º, 10 e 11 da Lei de Improbidade e definiu regras de transição. A revogação da modalidade culposa não desfaz coisa julgada; pode alcançar fatos culposos ainda sem condenação definitiva, com análise de eventual dolo. O novo regime prescricional não retroage.')

    h(doc, '6 Estabilidade', 1)
    h(doc, '6.1 Aquisição e perda', 2)
    p(doc, 'Estabilidade não é sinônimo de efetividade nem de estágio probatório. Efetividade descreve a forma de provimento do cargo. Estágio probatório é avaliação. Estabilidade é garantia constitucional adquirida após três anos de efetivo exercício e avaliação especial por comissão, para quem foi nomeado em cargo efetivo por concurso.')
    table(doc, ['Hipótese', 'Efeito possível'], [
        ['Sentença judicial transitada em julgado', 'Perda do cargo'],
        ['Processo administrativo com ampla defesa', 'Perda do cargo'],
        ['Avaliação periódica de desempenho', 'Perda na forma de lei complementar com ampla defesa'],
        ['Excesso de despesa com pessoal', 'Perda excepcional após as medidas prévias do art. 169'],
    ])

    h(doc, '6.2 Reintegração recondução e disponibilidade', 2)
    table(doc, ['Instituto', 'Quando aparece'], [
        ['Reintegração', 'Retorno do servidor estável cuja demissão foi invalidada'],
        ['Recondução', 'Retorno do ocupante estável ao cargo de origem nas hipóteses pertinentes'],
        ['Aproveitamento', 'Retorno do servidor em disponibilidade a cargo adequado'],
        ['Disponibilidade', 'Situação temporária com remuneração proporcional até aproveitamento'],
    ])
    p(doc, 'Se a demissão for invalidada, o servidor é reintegrado. O ocupante da vaga, se estável, é reconduzido sem indenização, aproveitado em outro cargo ou posto em disponibilidade. Se o cargo for extinto ou declarado desnecessário, o estável fica em disponibilidade até o adequado aproveitamento.')

    h(doc, '7 Previdência dos servidores', 1)
    h(doc, '7.1 Quem pertence ao RPPS', 2)
    p(doc, 'O RPPS do art. 40 protege titulares de cargos efetivos. É contributivo e solidário, com participação do ente, dos ativos, aposentados e pensionistas e com preservação do equilíbrio financeiro e atuarial. Ocupante exclusivamente de cargo em comissão, empregado público e temporário pertencem em regra ao RGPS.')
    table(doc, ['Tema', 'Regra central'], [
        ['Aposentadoria compulsória', 'Setenta e cinco anos para os agentes alcançados pela disciplina constitucional e complementar'],
        ['Aposentadoria voluntária da União', 'Idades mínimas constitucionais de 62 anos para mulher e 65 para homem além dos demais requisitos'],
        ['Estados DF e Municípios', 'Requisitos dependem da Constituição e das reformas do respectivo ente'],
        ['Previdência complementar', 'Limita benefícios do RPPS ao teto do RGPS para os alcançados pelo novo regime'],
        ['Abono de permanência', 'Pode ser instituído para quem completou requisitos e permanece em atividade'],
        ['Novos regimes próprios', 'A EC 103/2019 vedou a criação de novos RPPS'],
    ])
    p(doc, 'Erro comum. Copiar automaticamente para Estados e Municípios todos os requisitos aplicáveis à União. Após a EC 103/2019, parte importante das regras depende da reforma de cada ente.')

    h(doc, '7.2 Atualização da EC 140 de 2026', 2)
    p(doc, 'A EC 140/2026 acrescentou tratamento previdenciário diferenciado para agentes comunitários de saúde e agentes de combate às endemias. A regra permanente do art. 40, § 5º-A exige idade de 57 anos para mulher e 60 para homem, com 25 anos de contribuição e de efetivo exercício na atividade. A própria emenda também criou regras de transição e alterou os arts. 198 e 201.')
    p(doc, 'Como estudar. Primeiro identifique se o caso é RPPS ou RGPS. Depois verifique se a questão cobra regra permanente ou transição. Só então aplique idade, tempo de contribuição e tempo na atividade.')

    h(doc, '8 Jurisprudência e atualizações em mapa', 1)
    table(doc, ['Referência', 'Tese para a prova'], [
        ['ADI 2.135', 'Validade da EC 19/1998 quanto à retirada da obrigatoriedade constitucional do regime jurídico único'],
        ['Tema 161', 'Direito à nomeação de aprovado dentro das vagas como regra'],
        ['Tema 1.164', 'Mitigação excepcional diante da extinção motivada de cargos por limite prudencial antes do fim da validade'],
        ['Tema 531', 'Desconto dos dias de greve em regra com compensação possível e exceção por ilícito estatal'],
        ['Tema 1.199', 'Dolo nos arts. 9º 10 e 11 da LIA e regras específicas de transição'],
        ['Súmula Vinculante 13', 'Vedação ao nepotismo e aos ajustes recíprocos'],
        ['EC 138/2025', 'Professor pode acumular com outro cargo de qualquer natureza se houver compatibilidade'],
        ['EC 140/2026', 'Aposentadoria diferenciada de agentes comunitários de saúde e de combate às endemias'],
    ])

    h(doc, '9 Dez questões de aplicação', 1)
    p(doc, 'Resolva antes de consultar o comentário. As quatro primeiras questões são fáceis, as quatro seguintes são médias e as duas últimas são difíceis.')
    for q in QUESTIONS:
        add_question(doc, q)

    h(doc, '10 Conteúdo resumido', 1)
    p(doc, 'Estrutura. Sentido subjetivo pergunta quem atua; sentido objetivo pergunta o que é feito. Órgãos pertencem à Administração direta e não têm personalidade própria; entidades da indireta possuem personalidade.')
    p(doc, 'Princípios. LIMPE significa legalidade, impessoalidade, moralidade, publicidade e eficiência. Publicidade institucional não pode virar promoção pessoal.')
    p(doc, 'Ingresso. Cargo e emprego exigem concurso; cargo em comissão é exceção. Função de confiança é exclusiva de efetivo. Temporário exige lei, prazo e necessidade temporária excepcional.')
    p(doc, 'Servidores. Remuneração depende de lei; teto e subtetos variam. Acumulação exige hipótese constitucional e horários compatíveis. Estabilidade pertence ao concursado em cargo efetivo após três anos e avaliação.')
    p(doc, 'Responsabilidade e previdência. Estado responde objetivamente perante a vítima e cobra regressivamente do agente com dolo ou culpa. RPPS alcança cargo efetivo; comissão exclusiva, temporário e emprego público seguem em regra o RGPS.')

    h(doc, '11 Cinco questões de revisão', 1)
    p(doc, 'Julgue certo ou errado e explique o motivo antes de consultar a resposta.')
    for review_index, q in enumerate(REVIEW_QUESTIONS):
        if review_index == 2:
            doc.add_page_break()
        add_question(doc, q, review=True)

    h(doc, '12 Revisão ativa', 1)
    h(doc, 'Revisão de cinco minutos', 2)
    bullets(doc, [
        'Explique a diferença entre sentido subjetivo e objetivo.',
        'Recite o LIMPE e dê um exemplo de violação de cada princípio.',
        'Compare cargo efetivo emprego função de confiança comissão e temporário.',
        'Reconstrua prazo prorrogação prioridade e direito à nomeação.',
        'Liste as três acumulações permitidas após a EC 138/2025.',
        'Explique responsabilidade objetiva e direito de regresso.',
        'Diferencie efetividade estágio probatório e estabilidade.',
        'Diga quem pertence ao RPPS e qual novidade trouxe a EC 140/2026.',
    ])

    h(doc, 'O módulo em um minuto', 2)
    p(doc, 'Toda atuação administrativa começa pelo art. 37 e pelo LIMPE. O ingresso normal ocorre por concurso; confiança, comissão e temporário possuem limites próprios. Remuneração, teto e acumulação dependem do texto constitucional atualizado. Dano causado pelo agente gera responsabilidade objetiva do Estado e regresso com dolo ou culpa. Estabilidade protege apenas o concursado em cargo efetivo que cumpre os requisitos do art. 41. O RPPS pertence aos titulares de cargo efetivo e deve ser estudado com atenção às reformas do respectivo ente.')

    h(doc, 'O que decorar', 2)
    bullets(doc, [
        'LIMPE: legalidade impessoalidade moralidade publicidade eficiência.',
        'Concurso: até dois anos e uma prorrogação por igual período.',
        'Função de confiança: somente efetivo.',
        'Comissão: direção chefia e assessoramento.',
        'Temporário: necessidade temporária de excepcional interesse público.',
        'Acumulação: dois professores; professor mais qualquer outro; dois vínculos privativos de saúde.',
        'Improbidade: suspensão de direitos políticos e dolo nos arts. 9º 10 e 11.',
        'Estabilidade: cargo efetivo concurso três anos e avaliação especial.',
        'Responsabilidade: objetiva para o Estado e subjetiva no regresso.',
        'RPPS: titular de cargo efetivo.',
    ])

    h(doc, 'Pegadinhas que a prova usa', 2)
    p(doc, '“Todo agente público adquire estabilidade.” Errado: a garantia do art. 41 exige cargo efetivo, concurso, três anos e avaliação especial.')
    p(doc, '“Cargo em comissão serve para qualquer atividade.” Errado: destina-se somente a direção, chefia e assessoramento.')
    p(doc, '“Professor só pode acumular com cargo técnico ou científico.” Desatualizado: a EC 138/2025 passou a permitir outro cargo de qualquer natureza, com horários compatíveis.')
    p(doc, '“A vítima precisa provar culpa do agente.” Errado para a ação contra o Estado ou prestador; dolo ou culpa aparecem no regresso.')
    p(doc, '“Improbidade gera perda dos direitos políticos.” Errado: a Constituição prevê suspensão.')
    p(doc, '“A cautelar da ADI 2.135 ainda é a última palavra.” Errado: o mérito foi concluído em 2024 e validou a alteração da EC 19/1998.')

    h(doc, 'Erros de estudo a corrigir', 2)
    p(doc, 'Decorar o LIMPE sem saber aplicá-lo. Correção: associe cada princípio a um caso de concurso, publicidade, nomeação ou gasto público.')
    p(doc, 'Misturar vínculo com função. Correção: primeiro identifique se há cargo, emprego, comissão, função de confiança ou contrato temporário.')
    p(doc, 'Usar jurisprudência antiga sem conferir a data. Correção: destaque ADI 2.135, Tema 1.164 e as ECs 138 e 140 nas revisões.')
    p(doc, 'Estudar aposentadoria sem identificar o ente e o regime. Correção: separe União, regras locais, RPPS, RGPS, regra permanente e transição.')

    h(doc, 'Checklist de domínio', 2)
    for item in [
        '☐ Distingo Administração em sentido subjetivo e objetivo.',
        '☐ Explico Administração direta e indireta sem confundir órgão e entidade.',
        '☐ Aplico cada princípio do LIMPE a um caso concreto.',
        '☐ Comparo cargo emprego função de confiança comissão e temporário.',
        '☐ Resolvo prazo prorrogação prioridade e direito à nomeação em concursos.',
        '☐ Explico nepotismo greve remuneração teto e subsídio.',
        '☐ Aplico as acumulações permitidas após a EC 138/2025.',
        '☐ Distingo responsabilidade objetiva do Estado e regresso contra o agente.',
        '☐ Explico improbidade estabilidade reintegração recondução e disponibilidade.',
        '☐ Identifico RPPS RGPS e a atualização da EC 140/2026.',
        '☐ Justifico as respostas das quinze questões do módulo.',
    ]:
        p(doc, item)
    p(doc, 'Se algum item ainda depender de adivinhação, retorne à aula correspondente antes de marcar o módulo como concluído.')

    h(doc, '13 Referências e fontes de consulta', 1)
    p(doc, 'Constituição da República Federativa do Brasil de 1988. Texto compilado, especialmente arts. 37 a 41. Consulta em 22 de outubro de 2026.')
    p(doc, 'Fonte oficial: https://www.planalto.gov.br/ccivil_03/constituicao/constituicaocompilado.htm')
    p(doc, 'Emenda Constitucional 138 de 2025. Acumulação de cargo de professor com outro de qualquer natureza.')
    p(doc, 'Fonte oficial: https://www.planalto.gov.br/ccivil_03/constituicao/emendas/emc/emc138.htm')
    p(doc, 'Emenda Constitucional 140 de 2026. Aposentadoria diferenciada de agentes comunitários de saúde e agentes de combate às endemias.')
    p(doc, 'Fonte oficial: https://www.planalto.gov.br/ccivil_03/constituicao/emendas/emc/emc140.htm')
    p(doc, 'Supremo Tribunal Federal. ADI 2.135, Tema 1.164, Tema 531, Tema 1.199 e Súmula Vinculante 13.')
    p(doc, 'Fonte oficial: https://portal.stf.jus.br/')
    p(doc, 'Base de partida: arquivo CF M05 de teoria completa enviado pelo usuário. O conteúdo foi reorganizado pedagogicamente e atualizado quanto ao art. 39 e às Emendas Constitucionais 138 e 140.')


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    doc = Document(str(TEMPLATE))
    B.clear_body(doc)
    doc.core_properties.title = 'Administração Pública e servidores públicos'
    doc.core_properties.subject = 'CF M05 Apostila didática'
    doc.core_properties.keywords = 'Direito Constitucional, Administração Pública, servidores públicos, CF M05'
    update_running_text(doc)
    add_cover(doc)
    build_body(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(json.dumps({'output': str(OUT), 'bytes': OUT.stat().st_size, 'questions': 15}, ensure_ascii=False))


if __name__ == '__main__':
    main()
